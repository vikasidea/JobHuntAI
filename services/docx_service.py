from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_resume_docx(resume_text, output_path):

    document = Document()

    # Default font
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    lines = resume_text.splitlines()

    section_headers = {
        "NAME AND CONTACT INFORMATION",
        "PROFESSIONAL SUMMARY",
        "SKILLS",
        "EDUCATION",
        "INTERNSHIPS / EXPERIENCE",
        "PROJECTS",
        "CERTIFICATIONS",
        "ACHIEVEMENTS / EXTRACURRICULAR ACTIVITIES"
    }

    for line in lines:

        line = line.strip()

        if not line:
            continue

        # Ignore markdown separators
        if line == "---":
            continue

        # Section heading
        if line.upper() in section_headers:

            paragraph = document.add_paragraph()
            run = paragraph.add_run(line.upper())
            run.bold = True
            run.font.size = Pt(11)

            continue

        # Bullet point
        if line.startswith("- "):

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                line[2:].strip()
            )

            continue

        # Normal line
        paragraph = document.add_paragraph()

        run = paragraph.add_run(line)

        # Make first line/name slightly larger
        if line.upper() == "VIKAS YADAV":

            run.bold = True
            run.font.size = Pt(16)

            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.save(output_path)